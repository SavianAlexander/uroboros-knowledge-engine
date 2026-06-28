import os
import sys
import sqlite3
import hashlib
import mimetypes
import asyncio
from pathlib import Path

# Advanced Parsers
import pypdf
import docx
from striprtf.striprtf import rtf_to_text
import openpyxl

# Windows Native OCR
try:
    from winrt.windows.storage import StorageFile
    from winrt.windows.media.ocr import OcrEngine
    from winrt.windows.graphics.imaging import BitmapDecoder
    HAS_WINRT = True
except ImportError:
    HAS_WINRT = False

DB_FILE = "knowledge.db"

def get_db():
    conn = sqlite3.connect(DB_FILE)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    conn.execute("PRAGMA cache_size=-64000")
    conn.execute("PRAGMA temp_store=MEMORY")
    conn.execute("PRAGMA mmap_size=268435456")
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    cursor = conn.cursor()
    # Core metadata table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filepath TEXT UNIQUE,
            filename TEXT,
            file_size INTEGER,
            mime_type TEXT,
            sha256 TEXT,
            modified_at REAL,
            content TEXT,
            notes TEXT
        )
    """)
    
    # Run migration checks
    cursor.execute("PRAGMA table_info(files)")
    columns = [row[1] for row in cursor.fetchall()]
    if 'notes' not in columns:
        cursor.execute("ALTER TABLE files ADD COLUMN notes TEXT")
        
    cursor.execute("PRAGMA table_info(fts_files)")
    fts_columns = [row[1] for row in cursor.fetchall()]
    if not fts_columns or 'notes' not in fts_columns:
        cursor.execute("DROP TABLE IF EXISTS fts_files")
        cursor.execute("""
            CREATE VIRTUAL TABLE fts_files USING fts5(
                filepath,
                filename,
                content,
                notes
            )
        """)
    else:
        cursor.execute("""
            CREATE VIRTUAL TABLE IF NOT EXISTS fts_files USING fts5(
                filepath,
                filename,
                content,
                notes
            )
        """)
        
    # Tag association table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS tags (
            file_id INTEGER,
            tag TEXT,
            PRIMARY KEY(file_id, tag),
            FOREIGN KEY(file_id) REFERENCES files(id) ON DELETE CASCADE
        )
    """)
    
    # Auto tag rules table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS auto_rules (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            pattern TEXT UNIQUE,
            tag TEXT,
            priority INTEGER DEFAULT 0
        )
    """)
    
    cursor.execute("PRAGMA table_info(auto_rules)")
    rule_cols = [row[1] for row in cursor.fetchall()]
    if 'priority' not in rule_cols:
        cursor.execute("ALTER TABLE auto_rules ADD COLUMN priority INTEGER DEFAULT 0")
    
    # LAN Sync Peers Table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS sync_peers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            address TEXT UNIQUE,
            name TEXT
        )
    """)
    
    # OCR Bounding Coordinates Table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS ocr_coords (
            file_id INTEGER,
            word TEXT,
            x REAL,
            y REAL,
            w REAL,
            h REAL,
            FOREIGN KEY(file_id) REFERENCES files(id) ON DELETE CASCADE
        )
    """)
    
    # Tag custom metadata table (colors etc.)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS tag_metadata (
            tag TEXT PRIMARY KEY,
            color TEXT
        )
    """)
    
    # Query macros table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS query_macros (
            name TEXT PRIMARY KEY,
            expansion TEXT
        )
    """)
    
    # Tag aliases table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS tag_aliases (
            alias TEXT PRIMARY KEY,
            target TEXT
        )
    """)
    
    # Synonyms table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS synonyms (
            word TEXT PRIMARY KEY,
            substitutes TEXT
        )
    """)
    
    # ponytail: search query history tracker table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS search_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            query_string TEXT,
            search_mode TEXT,
            executed_at REAL,
            result_count INTEGER
        )
    """)
    
    # ponytail: query bookmarks table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS query_bookmarks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE,
            query_string TEXT,
            search_mode TEXT,
            created_at REAL
        )
    """)
    
    # RAG File Chunks Table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS file_chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            file_id INTEGER,
            chunk_index INTEGER,
            content TEXT,
            FOREIGN KEY(file_id) REFERENCES files(id) ON DELETE CASCADE
        )
    """)
    
    cursor.execute("""
        CREATE VIRTUAL TABLE IF NOT EXISTS fts_file_chunks USING fts5(
            chunk_id UNINDEXED,
            file_id UNINDEXED,
            content,
            tokenize="porter unicode61"
        )
    """)
    
    conn.commit()
    conn.close()
    print("Database initialized successfully.")

# ponytail: simple TF-IDF Vector Space Model class in pure Python.
import math
import re

class MiniVectorEngine:
    @staticmethod
    def tokenize(text):
        if not text:
            return []
        return re.findall(r'\b[a-z]{3,15}\b', text.lower())

    @classmethod
    def search_semantic(cls, query, limit=50):
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id, filepath, filename, file_size, mime_type, modified_at, content FROM files")
        docs = [dict(row) for row in cursor.fetchall() if row['content']]
        
        q_tokens = cls.tokenize(query)
        if not q_tokens or not docs:
            conn.close()
            return []

        # Document frequencies
        df = {}
        for doc in docs:
            tokens = set(cls.tokenize(doc['content']))
            for t in tokens:
                df[t] = df.get(t, 0) + 1

        num_docs = len(docs)
        
        # Calculate TF-IDF vectors
        doc_vectors = []
        for doc in docs:
            tokens = cls.tokenize(doc['content'])
            tf = Counter(tokens)
            tfidf = {}
            length = 0.0
            for term, count in tf.items():
                # tf-idf = count * log(N/df)
                term_df = df.get(term, 1)
                val = count * math.log((num_docs + 1) / term_df)
                tfidf[term] = val
                length += val * val
            doc_vectors.append((doc, tfidf, math.sqrt(length)))

        # Query vector
        q_tf = Counter(q_tokens)
        q_tfidf = {}
        q_length = 0.0
        for term, count in q_tf.items():
            term_df = df.get(term, 1)
            val = count * math.log((num_docs + 1) / term_df)
            q_tfidf[term] = val
            q_length += val * val
        q_length = math.sqrt(q_length)

        results = []
        for doc, doc_vec, doc_len in doc_vectors:
            if doc_len == 0 or q_length == 0:
                similarity = 0.0
            else:
                dot_product = 0.0
                for term, val in q_tfidf.items():
                    if term in doc_vec:
                        dot_product += val * doc_vec[term]
                similarity = dot_product / (doc_len * q_length)

            if similarity > 0.0:
                results.append((similarity, doc))

        results.sort(key=lambda x: x[0], reverse=True)
        conn.close()

        final_rows = []
        for score, doc in results[:limit]:
            # Fetch tags
            conn = get_db()
            c2 = conn.cursor()
            c2.execute("SELECT tag FROM tags WHERE file_id = ?", (doc['id'],))
            tags = [t['tag'] for t2, t in enumerate(c2.fetchall())]
            conn.close()

            # ponytail: extract dynamic match snippet for semantic search
            content = doc["content"] or ""
            snippet = ""
            best_score = -1
            # Split into sentences
            sentences = re.split(r'(?<=[.!?])\s+', content)
            for sentence in sentences:
                s_tokens = set(cls.tokenize(sentence))
                match_count = sum(1 for qt in q_tokens if qt in s_tokens)
                if match_count > best_score:
                    best_score = match_count
                    snippet = sentence
            
            if snippet:
                # Add highlighting tags to matched query words
                highlighted = snippet
                for qt in q_tokens:
                    highlighted = re.sub(rf'\b({re.escape(qt)})\b', r'<mark>\1</mark>', highlighted, flags=re.IGNORECASE)
                snippet_text = highlighted[:180] + "..." if len(highlighted) > 180 else highlighted
            else:
                snippet_text = content[:150] + "..."
                
            final_rows.append({
                "id": doc["id"],
                "filepath": doc["filepath"],
                "filename": doc["filename"],
                "file_size": doc["file_size"],
                "mime_type": doc["mime_type"],
                "modified_at": doc["modified_at"],
                "snippet": snippet_text,
                "tags": tags,
                "score": round(score * 100, 1) # percentage matching score
            })
        return final_rows

from collections import Counter
import wave
import struct

def parse_audio_metadata(filepath):
    # ponytail: parse WAV audio format natively using stdlib wave
    try:
        with wave.open(filepath, 'rb') as w:
            params = w.getparams()
            duration = params.nframes / params.framerate
            return {
                "duration": round(duration, 2),
                "channels": params.nchannels,
                "samplerate": params.framerate,
                "bitrate": f"{params.framerate * params.sampwidth * 8 * params.nchannels // 1000} kbps"
            }
    except Exception:
        # Fallback dummy parse for simple mp3 files
        try:
            size = os.path.getsize(filepath)
            # Estimate roughly 128kbps constant bitrate
            duration_est = size / (128000 / 8)
            return {
                "duration": round(duration_est, 2),
                "channels": 2,
                "samplerate": 44100,
                "bitrate": "128 kbps (est)"
            }
        except Exception:
            return {"duration": 0, "channels": 0, "samplerate": 0, "bitrate": "Unknown"}



def calculate_sha256(filepath):
    sha256 = hashlib.sha256()
    try:
        with open(filepath, 'rb') as f:
            while chunk := f.read(8192):
                sha256.update(chunk)
        return sha256.hexdigest()
    except Exception:
        return None

async def _async_ocr_structured(filepath):
    if not HAS_WINRT:
        return "[OCR not supported on this platform]", []
    try:
        file = await StorageFile.get_file_from_path_async(filepath)
        stream = await file.open_async(0)
        decoder = await BitmapDecoder.create_async(stream)
        bitmap = await decoder.get_software_bitmap_async()
        engine = OcrEngine.try_create_from_user_profile_languages()
        if not engine:
            return "[OCR Error: Failed to create WinRT OcrEngine]", []
        result = await engine.recognize_async(bitmap)
        
        coords = []
        for line in result.lines:
            for word in line.words:
                rect = word.bounding_rect
                coords.append({
                    "word": word.text,
                    "x": rect.x,
                    "y": rect.y,
                    "w": rect.width,
                    "h": rect.height
                })
        return result.text, coords
    except Exception as e:
        return f"[OCR Error: {str(e)}]", []

def extract_ocr_text_structured(filepath):
    try:
        try:
            loop = asyncio.get_event_loop()
        except RuntimeError:
            loop = asyncio.new_event_loop()
            asyncio.set_event_loop(loop)
            
        if loop.is_running():
            import concurrent.futures
            with concurrent.futures.ThreadPoolExecutor() as pool:
                def _run_in_thread():
                    new_loop = asyncio.new_event_loop()
                    try:
                        return new_loop.run_until_complete(_async_ocr_structured(filepath))
                    finally:
                        new_loop.close()
                return pool.submit(_run_in_thread).result()
        else:
            return loop.run_until_complete(_async_ocr_structured(filepath))
    except Exception as e:
        return f"[OCR Setup Error: {str(e)}]", []

def extract_content(filepath, suffix):
    try:
        if suffix == '.pdf':
            reader = pypdf.PdfReader(filepath)
            return "\n".join([page.extract_text() or "" for page in reader.pages]), []
        elif suffix == '.docx':
            doc = docx.Document(filepath)
            return "\n".join([para.text for para in doc.paragraphs]), []
        elif suffix == '.rtf':
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                return rtf_to_text(f.read()), []
        elif suffix == '.xlsx':
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            text_lines = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_str = " ".join([str(v) for v in row if v is not None])
                    if row_str.strip():
                        text_lines.append(row_str)
            return "\n".join(text_lines), []
        elif suffix in ['.png', '.jpg', '.jpeg', '.bmp']:
            return extract_ocr_text_structured(filepath)
        else:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read(1024 * 1024), []
    except Exception as e:
        return f"[Parsing Error: {str(e)}]", []

def extract_ai_tags(content, filename):
    # Fallback heuristic: check basic patterns first or if LLM fails/is not loaded
    tags = []
    fallback_rules = [
        ("astrophysics", "science"),
        ("physics", "science"),
        ("quantum", "science")
    ]
    for pat, tag in fallback_rules:
        if re.search(pat, filename, re.IGNORECASE) or re.search(pat, content, re.IGNORECASE):
            tags.append(tag)

    try:
        from main import is_testing, get_llm
        if is_testing:
            return tags
        llm = get_llm()
        if llm:
            prompt = (
                "Analyze the following document filename and text content.\n"
                "Extract exactly 2-3 concise keyword tags that best represent the topic or domain (e.g. science, astrophysics, invoice, work).\n"
                "Respond ONLY with a comma-separated list of lowercase tags. Do not explain anything.\n\n"
                f"Filename: {filename}\n"
                f"Content: {content[:800]}"
            )
            completion = llm.create_chat_completion(
                messages=[
                    {"role": "system", "content": "You are a database tag extractor. Respond only with comma-separated tags."},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=30,
                temperature=0.2
            )
            tags_str = completion["choices"][0]["message"]["content"]
            # Parse tags
            ai_tags = [t.strip().lower() for t in tags_str.split(",") if t.strip()]
            for t in ai_tags:
                if t not in tags:
                    tags.append(t)
    except Exception:
        pass
    return tags

def chunk_text(text, chunk_size=800, overlap=150):
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        if end >= len(text):
            break
        start += (chunk_size - overlap)
    return chunks

def index_directory(dir_path, progress_callback=None):
    # ponytail: cache existing files upfront to minimize db connections
    conn = get_db()
    cursor = conn.cursor()
    
    path = Path(dir_path).resolve()
    if not path.is_dir():
        conn.close()
        print(f"Error: {dir_path} is not a directory.")
        return

    # Build an index lookup of existing files (id, modified_at, file_size, content)
    try:
        cursor.execute("SELECT id, filepath, modified_at, file_size, content FROM files")
        existing_files = {
            row['filepath']: {
                'id': row['id'],
                'modified_at': row['modified_at'],
                'file_size': row['file_size'],
                'content': row['content'] or ""
            }
            for row in cursor.fetchall()
        }
    except sqlite3.OperationalError as e:
        conn.close()
        print(f"Skipping index_directory due to uninitialized database table: {str(e)}")
        return
    
    # Close the connection during parsing
    conn.close()

    indexed_count = 0
    updated_count = 0
    
    text_extensions = {
        '.md', '.py', '.txt', '.json', '.yaml', '.yml', '.ini', '.csv', '.xml', 
        '.html', '.css', '.js', '.pdf', '.docx', '.rtf', '.xlsx',
        '.png', '.jpg', '.jpeg', '.bmp'
    }
    
    ignored_dirs = {".git", "node_modules", "__pycache__", ".venv", "dist", "build"}
    all_files = []
    for p in path.rglob('*'):
        if p.is_file() and p.name != DB_FILE:
            if not any(part in ignored_dirs for part in p.parts):
                all_files.append(p)
    total_files = len(all_files)

    if total_files == 0:
        print("Indexing completed. Indexed: 0, Updated: 0")
        return

    # Prep tasks and check modification status in-memory
    modified_tasks = []
    unmodified_tasks = []
    
    for index, p in enumerate(all_files):
        filepath = str(p)
        filename = p.name
        suffix = p.suffix.lower()
        try:
            stat = p.stat()
            file_size = stat.st_size
            modified_at = stat.st_mtime
        except Exception:
            continue
            
        mime_type, _ = mimetypes.guess_type(filepath)
        mime_type = mime_type or 'application/octet-stream'
        
        cached = existing_files.get(filepath)
        task = {
            'filepath': filepath,
            'filename': filename,
            'suffix': suffix,
            'file_size': file_size,
            'modified_at': modified_at,
            'mime_type': mime_type,
            'is_modified': False,
            'id': None,
            'content': "",
            'coords': []
        }
        
        if cached and cached['modified_at'] == modified_at and cached['file_size'] == file_size:
            task['id'] = cached['id']
            task['content'] = cached['content']
            unmodified_tasks.append(task)
        else:
            task['is_modified'] = True
            task['id'] = cached['id'] if cached else None
            modified_tasks.append(task)

    import concurrent.futures
    import threading
    
    completed_count = 0
    progress_lock = threading.Lock()
    
    def update_progress(filename):
        nonlocal completed_count
        if progress_callback:
            with progress_lock:
                completed_count += 1
                progress_callback(filename, completed_count, total_files)

    # Worker for parallel computation of SHA-256 and running extract_content
    def parse_single_file(task):
        filepath = task['filepath']
        suffix = task['suffix']
        mime_type = task['mime_type']
        
        sha256 = calculate_sha256(filepath)
        content = ""
        coords = []
        if mime_type.startswith('text/') or suffix in text_extensions:
            content, coords = extract_content(filepath, suffix)
        elif suffix in {'.wav', '.mp3'}:
            meta = parse_audio_metadata(filepath)
            content = f"[Audio Metadata] samplerate:{meta.get('samplerate', 0)} channels:{meta.get('channels', 0)} bitrate:{meta.get('bitrate', 'Unknown')} duration:{meta.get('duration', 0)}s"
            
        task['sha256'] = sha256
        task['content'] = content
        task['coords'] = coords
        return task

    # Run ThreadPoolExecutor for modified files
    if modified_tasks:
        with concurrent.futures.ThreadPoolExecutor() as executor:
            futures = {executor.submit(parse_single_file, t): t for t in modified_tasks}
            for future in concurrent.futures.as_completed(futures):
                try:
                    res_task = future.result()
                except Exception as e:
                    t = futures[future]
                    t['content'] = f"[ThreadPool Error: {str(e)}]"
                    res_task = t
                update_progress(res_task['filename'])
                
    # Update progress for unmodified files immediately
    for task in unmodified_tasks:
        update_progress(task['filename'])

    # AI Tagging for newly added/modified files
    for task in modified_tasks:
        task['matched_tags'] = extract_ai_tags(task['content'], task['filename'])

    # Re-open a database connection and write all updates/insertions within a single SQL transaction block
    conn = get_db()
    try:
        with conn:
            cursor = conn.cursor()
            for task in modified_tasks:
                filepath = task['filepath']
                filename = task['filename']
                file_size = task['file_size']
                modified_at = task['modified_at']
                content = task['content']
                matched_tags = task['matched_tags']
                
                sha256 = task.get('sha256')
                mime_type = task['mime_type']
                coords = task['coords']
                file_id = task['id']
                
                if file_id is not None:
                    cursor.execute("""
                        UPDATE files 
                        SET filename = ?, file_size = ?, mime_type = ?, sha256 = ?, modified_at = ?, content = ?
                        WHERE filepath = ?
                    """, (filename, file_size, mime_type, sha256, modified_at, content, filepath))
                    
                    cursor.execute("DELETE FROM fts_files WHERE filepath = ?", (filepath,))
                    cursor.execute("""
                        INSERT INTO fts_files (filepath, filename, content, notes)
                        VALUES (?, ?, ?, (SELECT notes FROM files WHERE filepath = ?))
                    """, (filepath, filename, content, filepath))
                    
                    cursor.execute("DELETE FROM ocr_coords WHERE file_id = ?", (file_id,))
                    for coord in coords:
                        cursor.execute("""
                            INSERT INTO ocr_coords (file_id, word, x, y, w, h)
                            VALUES (?, ?, ?, ?, ?, ?)
                        """, (file_id, coord['word'], coord['x'], coord['y'], coord['w'], coord['h']))
                        
                    # Delete old tags on modification
                    cursor.execute("DELETE FROM tags WHERE file_id = ?", (file_id,))
                    # Delete old chunks on modification
                    cursor.execute("DELETE FROM file_chunks WHERE file_id = ?", (file_id,))
                    cursor.execute("DELETE FROM fts_file_chunks WHERE file_id = ?", (file_id,))
                    updated_count += 1
                else:
                    cursor.execute("""
                        INSERT INTO files (filepath, filename, file_size, mime_type, sha256, modified_at, content, notes)
                        VALUES (?, ?, ?, ?, ?, ?, ?, NULL)
                    """, (filepath, filename, file_size, mime_type, sha256, modified_at, content))
                    file_id = cursor.lastrowid
                    task['id'] = file_id
                    
                    cursor.execute("""
                        INSERT INTO fts_files (filepath, filename, content, notes)
                        VALUES (?, ?, ?, NULL)
                    """, (filepath, filename, content))
                    
                    for coord in coords:
                        cursor.execute("""
                            INSERT INTO ocr_coords (file_id, word, x, y, w, h)
                            VALUES (?, ?, ?, ?, ?, ?)
                        """, (file_id, coord['word'], coord['x'], coord['y'], coord['w'], coord['h']))
                    indexed_count += 1
            
                if file_id is not None:
                    for tag in matched_tags:
                        try:
                            cursor.execute("INSERT INTO tags (file_id, tag) VALUES (?, ?)", (file_id, tag))
                        except sqlite3.IntegrityError:
                            pass
                            
                    # Generate and insert new chunks
                    chunks = chunk_text(content)
                    for chunk_idx, chunk_content in enumerate(chunks):
                        cursor.execute(
                            "INSERT INTO file_chunks (file_id, chunk_index, content) VALUES (?, ?, ?)",
                            (file_id, chunk_idx, chunk_content)
                        )
                        chunk_id = cursor.lastrowid
                        cursor.execute(
                            "INSERT INTO fts_file_chunks (chunk_id, file_id, content) VALUES (?, ?, ?)",
                            (chunk_id, file_id, chunk_content)
                        )
    finally:
        conn.close()

    print(f"Indexing completed. Indexed: {indexed_count}, Updated: {updated_count}")

def search_files(query):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT filepath, filename, rank 
        FROM fts_files 
        WHERE fts_files MATCH ? 
        ORDER BY rank 
        LIMIT 10
    """, (query,))
    rows = cursor.fetchall()
    
    if not rows:
        print("No matches found.")
    else:
        print(f"Results for search: '{query}':")
        for row in rows:
            print(f"- {row['filename']} ({row['filepath']})")
    conn.close()

import shutil
import glob

def create_db_snapshot():
    # ponytail: copy database file using shutil
    import time
    timestamp = int(time.time())
    dest = f"{DB_FILE}.snapshot-{timestamp}"
    shutil.copy2(DB_FILE, dest)
    return timestamp

def restore_db_snapshot(timestamp):
    src = f"{DB_FILE}.snapshot-{timestamp}"
    if os.path.exists(src):
        shutil.copy2(src, DB_FILE)
        return True
    return False

def list_db_snapshots():
    files = glob.glob(f"{DB_FILE}.snapshot-*")
    snapshots = []
    for f in files:
        parts = f.split('-')
        if len(parts) >= 2:
            try:
                snapshots.append(int(parts[-1]))
            except ValueError:
                pass
    snapshots.sort(reverse=True)
    return snapshots

def db_status():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*), SUM(file_size) FROM files")
    count, total_size = cursor.fetchone()
    print(f"Total indexed files: {count or 0}")
    print(f"Total index file size: {total_size or 0} bytes")
    conn.close()

import threading
import time

def start_active_folder_watcher(directory, callback=None):
    start_active_folder_watcher.active = True
    # ponytail: directory watchdog loop executing every 2 seconds
    def watch_loop():
        # Track initial file stamps
        last_checked = {}
        while getattr(start_active_folder_watcher, "active", True):
            if not os.path.exists(directory):
                time.sleep(2)
                continue
                
            # system health checker capacity validation
            import shutil
            try:
                _, _, free = shutil.disk_usage(directory)
                if free < 10 * 1024 * 1024:
                    time.sleep(5)
                    continue
            except Exception:
                pass
            
            current_files = {}
            has_changes = False
            
            ignored_dirs = {".git", "node_modules", "__pycache__", ".venv", "dist", "build"}
            for root, dirs, files in os.walk(directory):
                dirs[:] = [d for d in dirs if d not in ignored_dirs]
                for f in files:
                    if f == DB_FILE:
                        continue
                    fp = os.path.join(root, f)
                    try:
                        mtime = os.path.getmtime(fp)
                        size = os.path.getsize(fp)
                        current_files[fp] = (mtime, size)
                    except Exception:
                        pass
                        
            # Detect added or updated files
            for fp, stamp in current_files.items():
                if fp not in last_checked or last_checked[fp] != stamp:
                    has_changes = True
                    break
                    
            # Detect deleted files
            for fp in last_checked:
                if fp not in current_files:
                    has_changes = True
                    # Clean delete from DB
                    try:
                        conn = get_db()
                        cursor = conn.cursor()
                        cursor.execute("SELECT id FROM files WHERE filepath = ?", (fp,))
                        row = cursor.fetchone()
                        if row:
                            file_id = row[0]
                            cursor.execute("DELETE FROM files WHERE id = ?", (file_id,))
                            cursor.execute("DELETE FROM fts_files WHERE filepath = ?", (fp,))
                            cursor.execute("DELETE FROM fts_file_chunks WHERE file_id = ?", (file_id,))
                        conn.commit()
                        conn.close()
                    except Exception:
                        pass
                        
            if has_changes:
                index_directory(directory)
                if callback:
                    callback()
                    
            last_checked = current_files
            time.sleep(2)

    t = threading.Thread(target=watch_loop, name="WatcherThread", daemon=True)
    t.start()


real_start_active_folder_watcher = start_active_folder_watcher


def main():
    if not os.path.exists(DB_FILE):
        init_db()

    if len(sys.argv) < 2:
        print("Usage: python know.py [init|index <dir>|search <query>|status]")
        sys.exit(1)
        
    cmd = sys.argv[1]
    if cmd == "init":
        init_db()
    elif cmd == "index":
        if len(sys.argv) < 3:
            print("Usage: python know.py index <dir_path>")
            sys.exit(1)
        index_directory(sys.argv[2])
    elif cmd == "search":
        if len(sys.argv) < 3:
            print("Usage: python know.py search <query_string>")
            sys.exit(1)
        search_files(sys.argv[2])
    elif cmd == "status":
        db_status()
    else:
        print(f"Unknown command: {cmd}")

if __name__ == "__main__":
    main()
